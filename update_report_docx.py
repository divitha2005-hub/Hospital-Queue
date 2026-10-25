import docx
from docx.shared import Pt
import os

def update_document(input_path, output_path):
    print(f"Loading document from {input_path}...")
    doc = docx.Document(input_path)
    
    # ---------------------------------------------------------
    # Helper functions
    # ---------------------------------------------------------
    def set_caption(p, prefix, text):
        p.text = ""
        r1 = p.add_run(prefix)
        r1.bold = True
        r1.font.size = Pt(10.5)
        r2 = p.add_run(text)
        r2.italic = True
        r2.font.size = Pt(10.5)

    def set_p(p, text):
        p.text = text

    def set_cell(cell, text, bold=False, italic=False, size_pt=10.0):
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size_pt)

    # ---------------------------------------------------------
    # 1. Update Media / Image Parts
    # ---------------------------------------------------------
    print("Updating embedded image parts...")
    asset_map = {
        'image12': 'new_report_assets/image12.png',
        'image13': 'new_report_assets/image13.png',
        'image14': 'new_report_assets/image14.png',
        'image15': 'new_report_assets/image15.png',
        'image16': 'new_report_assets/image16.png',
        'image17': 'new_report_assets/image17.png',
        'image28': 'new_report_assets/image28.png',
        'image29': 'new_report_assets/image29.png',
        'image30': 'new_report_assets/image30.png',
        'image31': 'new_report_assets/image31.png',
        'image32': 'new_report_assets/image32.png',
    }

    for rel in doc.part.rels.values():
        for key, asset_path in asset_map.items():
            if key in rel.target_ref:
                part = doc.part.related_parts[rel.rId]
                with open(asset_path, 'rb') as f:
                    part._blob = f.read()
                print(f"  Replaced {rel.target_ref} with {asset_path}")

    # ---------------------------------------------------------
    # 2. Section 1: Abstract & Metadata
    # ---------------------------------------------------------
    print("Updating Abstract and Metadata...")
    set_p(doc.paragraphs[22], 
        "The project uses three core design patterns that sit directly on its primary clinical use cases, "
        "and this report focuses strictly on those three: Factory Method (Creational), Facade (Structural), "
        "and State (Behavioural). PatientFactory (server/patterns/factory.js) turns the registration form into an "
        "EmergencyPatient, an OPDPatient, or a FollowUpPatient while enforcing domain-level category rules, such as "
        "auto-flooring emergency severity. AdmissionFacade (server/patterns/facade.js) provides the API layer with "
        "unified operations for admitting and discharging patients, coordinating bed allocation, doctor scheduling, "
        "and billing subsystems behind clean interfaces. The State pattern (server/patterns/state.js) governs "
        "patient lifecycle transitions (WaitingState -> ConsultationState -> AdmittedState -> DischargedState) "
        "and clinical escalation rules, ensuring that illegal transitions are blocked by transition guards and "
        "state-dependent operations are encapsulated in concrete state classes."
    )

    set_p(doc.paragraphs[23],
        "For every pattern, this report details the architectural intent, project justification, participating classes, "
        "class diagrams before and after applying the pattern, verified source code, and runtime verification from the "
        "application and test suites. We validated the entire clinical pathway through both the user interface and REST "
        "API endpoints. Automated test suites (npm test) verify 100% of the lifecycle state transitions, triage "
        "prioritization rules, and facade orchestration calls without regressions."
    )

    set_p(doc.paragraphs[24],
        "Keywords: design patterns, Factory Method, Facade, State pattern, hospital queue, triage, clinical lifecycle, Node.js, Express."
    )

    # ---------------------------------------------------------
    # 3. Section 2: Literature Survey
    # ---------------------------------------------------------
    print("Updating Literature Survey...")
    set_p(doc.paragraphs[27],
        "Design patterns are named, reusable solutions to recurring object-oriented software engineering challenges. "
        "The catalogue by Gamma, Helm, Johnson, and Vlissides describes 23 foundational patterns classified into creational, "
        "structural, and behavioural categories [1]. Our three patterns deliberately span these three categories: "
        "Factory Method is creational, Facade is structural, and State is behavioural. Each pattern addresses a distinct "
        "architectural concern in the hospital workflow: Factory Method governs heterogeneous patient instantiation based "
        "on triage category, Facade coordinates complex multi-subsystem hospital workflows (bed allocation, physician "
        "scheduling, invoice settlement) behind simplified interfaces, and State governs dynamic clinical lifecycle "
        "transitions and operational rules across a patient's care trajectory."
    )

    set_p(doc.paragraphs[34], "2.4 Clinical Lifecycle Governance and State Transitions")

    set_p(doc.paragraphs[35],
        "In healthcare systems, a patient progresses through distinct operational phases: waiting in the triage queue, "
        "active clinical consultation, inpatient ward admission, and final discharge. Each phase permits specific "
        "clinical and administrative actions while strictly forbidding others (e.g., discharging an outpatient before "
        "consultation or admitting an already discharged patient). Implementing this lifecycle through distributed string "
        "flags or conditional checks results in brittle, error-prone state corruption. The State pattern [1] encapsulates "
        "state-dependent behaviors into distinct classes derived from a common interface. State transitions are governed "
        "by strict transition guards, and each state manages its own operational rules (such as clinical triage escalation "
        "during consultation), keeping business logic cohesive and adherent to the Single Responsibility Principle."
    )

    # ---------------------------------------------------------
    # 4. Section 3: Research Gap
    # ---------------------------------------------------------
    print("Updating Research Gap and Objectives...")
    set_p(doc.paragraphs[40],
        "Most hospital queue demos concentrate on isolated utility tasks, such as token numbers or appointment booking, "
        "while generic pattern tutorials demonstrate concepts on toy problems (e.g., shapes or coffee). In contrast, "
        "BioCare EHR establishes an integrated domain where the three patterns collaborate directly: the domain entity "
        "created by PatientFactory enters the triage queue governed by WaitingState, transitions to ConsultationState "
        "upon physician call, is admitted and discharged across coordinated subsystems via AdmissionFacade, and "
        "transitions safely into terminal DischargedState governed by the State pattern."
    )

    set_p(doc.paragraphs[42],
        "If written naively, the application's complexity would collapse into monolithic route handlers: scattered "
        "instantiation conditionals, intertwined subsystem orchestrations (beds, doctors, billing) that risk incomplete "
        "transactions, and unvalidated status string assignments that allow illegal transitions. Introducing new patient "
        "types, modifying admission procedures, or refining clinical escalation paths would require high-risk modifications "
        "to central routing code."
    )

    set_p(doc.paragraphs[45],
        "Apply Factory Method, Facade, and State pattern at the three places in the system where they are architecturally "
        "justified, and demonstrate each in the production codebase."
    )

    # ---------------------------------------------------------
    # 5. Section 4: Design Patterns Overview (Table 1)
    # ---------------------------------------------------------
    print("Updating Section 4 Intro & Table 1...")
    set_p(doc.paragraphs[49],
        "The three patterns in this section are the main use case patterns of the project. Factory Method handles the "
        "registration of patients, Facade handles admission and discharge, and State handles patient lifecycle transitions "
        "and clinical governance. Each subsection follows the same order: intent, the need in our project, the classes "
        "involved, advantages, why we chose the pattern, class diagrams before and after, and the code. Screenshots of the "
        "running application for all three patterns are collected in Section 6."
    )

    # Table 0 (Table 1) Row 3:
    set_cell(doc.tables[0].rows[3].cells[0], "State", bold=True)
    set_cell(doc.tables[0].rows[3].cells[1], "Behavioural")
    set_cell(doc.tables[0].rows[3].cells[2], "Governing patient lifecycle transitions and clinical escalation")
    set_cell(doc.tables[0].rows[3].cells[3], "PatientState")
    set_cell(doc.tables[0].rows[3].cells[4], "server/patterns/state.js")

    # Table 3 (Facade Components) Row 1:
    set_cell(doc.tables[3].rows[1].cells[1], 'Single entry point. admitPatient(patient, bedCategory, durationDays) and dischargePatient(patient, insurancePlan).')

    # Update Facade discharge mention in P136:
    set_p(doc.paragraphs[136],
        "dischargePatient() in Figure 10 is the checkout. It verifies that the patient is currently in an admitted state, "
        "releases the bed, frees the supervising doctor, prepares the settled bill, and transitions the patient safely into "
        "DischargedState. The lines that change the status of the patient to discharged are fully validated through the State pattern."
    )

    # ---------------------------------------------------------
    # 6. Section 4.3: State Pattern Details
    # ---------------------------------------------------------
    print("Updating Section 4.3 (State Pattern)...")
    set_p(doc.paragraphs[139], "4.3 Behavioural Design Pattern: State")
    set_p(doc.paragraphs[140], "4.3.1 Intent")
    set_p(doc.paragraphs[141],
        "The State pattern allows an object to alter its behavior when its internal state changes, appearing as if the "
        "object changed its class [1]. In BioCare EHR, a patient transitions through defined clinical phases: WaitingState, "
        "ConsultationState, AdmittedState, and DischargedState. Rather than relying on scattered conditional logic and string "
        "comparisons throughout route handlers, patient behaviors (such as handling clinical escalation, bed allocation, "
        "or discharge) are delegated to the active PatientState object."
    )
    set_p(doc.paragraphs[142],
        "Flow: patient enters queue in WaitingState -> doctor calls patient triggering next(doctor) -> ConsultationState "
        "-> doctor admits patient triggering next('admit') -> AdmittedState -> checkout triggers next('discharge') -> "
        "DischargedState (terminal)."
    )

    set_p(doc.paragraphs[143], "4.3.2 Need in the Proposed System")
    set_p(doc.paragraphs[144],
        "A patient's healthcare lifecycle involves strict operational rules and sequencing constraints:\n"
        "1. A patient waiting in the queue cannot undergo ward admission or discharge without completing consultation.\n"
        "2. A consultation requires an assigned physician; attempting a consultation transition without doctor assignment must be blocked.\n"
        "3. During consultation, clinical escalation rules apply (e.g., severity >= 8 escalates to senior attending specialists with ICU priority, while mild symptoms are managed by junior staff).\n"
        "4. Admitted patients occupy physical hospital beds and require coordinated checkout before discharge.\n"
        "5. Discharged patients reach a terminal state where no further state transitions, consultations, or admissions are permitted."
    )
    set_p(doc.paragraphs[145],
        "Without the State pattern, enforcing these rules requires extensive, fragile if-else or switch statements across "
        "every controller. The State pattern encapsulates state-specific behaviors and transition guards cleanly inside dedicated classes."
    )

    set_p(doc.paragraphs[146], "4.3.3 Application Classes and Components")
    set_caption(doc.paragraphs[147], "Table 6: ", "Classes and components of the State pattern implementation")

    # Table 5: Components
    state_components = [
        ("PatientState", "Abstract base class declaring the clinical lifecycle interface: getStateName(), next(patient, option), and escalate(patient, issue). Throws descriptive domain errors by default for invalid operations."),
        ("WaitingState", "Initial queue state. Enforces doctor assignment guard when transitioning to ConsultationState. Strictly blocks premature ward admission or discharge."),
        ("ConsultationState", "Active physician evaluation. Executes clinical escalation rules based on severityScore and vitals. Transitions to AdmittedState (inpatient) or direct DischargedState (outpatient)."),
        ("AdmittedState", "Inpatient ward care. Tracks physical bed occupancy (ICU, Private Room, General Ward) and transitions strictly to DischargedState via checkout. Blocks queue re-entry or repeat consultations."),
        ("DischargedState", "Terminal state representing completed clinical visit. Rejects subsequent transition or escalation calls, guaranteeing post-discharge record immutability."),
        ("StateFactory", "Utility factory providing getState(stateName). Rehydrates concrete state class instances from persistent storage strings during database reads."),
        ("Patient (Context)", "Maintains current stateName and delegates next() transitions and escalate() requests directly to the active PatientState instance."),
        ("AdmissionFacade", "Coordinates subsystems during admission and discharge, validating that the patient is in the correct state before advancing to AdmittedState and DischargedState.")
    ]
    t5 = doc.tables[5]
    for idx, (comp, resp) in enumerate(state_components):
        set_cell(t5.rows[idx + 1].cells[0], comp, bold=True)
        set_cell(t5.rows[idx + 1].cells[1], resp)

    set_p(doc.paragraphs[149], "4.3.4 Advantages")
    set_p(doc.paragraphs[150], "Eliminates monolithic conditional logic: State-specific rules and transition validations are isolated from Express controllers.")
    set_p(doc.paragraphs[151], "Enforces robust transition guards: Illegal operations (e.g. discharging a waiting patient or transitioning from a discharged record) are caught and rejected immediately.")
    set_p(doc.paragraphs[152], "Encapsulates clinical escalation: Urgent triage escalation logic (routing severe cases to senior specialists) lives cohesively inside ConsultationState.")
    set_p(doc.paragraphs[153], "Maintains record integrity: DischargedState acts as a terminal guard, preventing accidental data modification on completed visits.")
    set_p(doc.paragraphs[154], "Supports clean state persistence: StateFactory seamlessly restores full object-oriented state instances from JSON/MongoDB records.")

    set_p(doc.paragraphs[155], "4.3.5 Why State Pattern was Selected")
    set_p(doc.paragraphs[156], "The patient lifecycle is inherently stateful with explicit transition rules, preconditions, and postconditions.")
    set_p(doc.paragraphs[157], "Clinical operations depend on current status: consultation actions differ fundamentally from triage waiting or ward admission.")
    set_p(doc.paragraphs[158], "Enforcing transition guards directly within the domain entity prevents data corruption across disparate API routes.")
    set_p(doc.paragraphs[159], "Separating state logic allows adding future states (such as ObservationState or ICUState) without modifying existing routing code.")

    set_p(doc.paragraphs[160], "4.3.6 Class Diagram Before Applying the Pattern")
    # P161 is image
    set_caption(doc.paragraphs[162], "Figure 11: ", "Structure before applying State Pattern: monolithic lifecycle with unvalidated state mutations and scattered conditional checks")

    set_p(doc.paragraphs[163], "4.3.7 Class Diagram After Applying the Pattern")
    # P164 is image
    set_caption(doc.paragraphs[165], "Figure 12: ", "Structure after applying State Pattern (PatientState, four concrete states, and StateFactory)")

    set_p(doc.paragraphs[166], "4.3.8 Code and Implementation")
    set_caption(doc.paragraphs[167], "Table 7: ", "Implementation summary for State Pattern")

    # Table 6: Implementation Summary
    t6 = doc.tables[6]
    set_cell(t6.rows[1].cells[0], "Source file", bold=True)
    set_cell(t6.rows[1].cells[1], "server/patterns/state.js")
    set_cell(t6.rows[2].cells[0], "Pattern role", bold=True)
    set_cell(t6.rows[2].cells[1], "State interface (PatientState), four concrete states (Waiting, Consultation, Admitted, Discharged), StateFactory for rehydration")
    set_cell(t6.rows[3].cells[0], "Context", bold=True)
    set_cell(t6.rows[3].cells[1], "Patient entity (server/models/patient.js), AdmissionFacade, and Express controllers in server/routes.js")
    set_cell(t6.rows[4].cells[0], "Runtime transitions", bold=True)
    set_cell(t6.rows[4].cells[1], "Waiting -> Consultation -> Admitted -> Discharged (Terminal)")

    # P169 is image
    set_caption(doc.paragraphs[170], "Figure 13: ", "Source code: PatientState, WaitingState and ConsultationState in server/patterns/state.js")
    # P171 is image
    set_caption(doc.paragraphs[172], "Figure 14: ", "Source code: AdmittedState, DischargedState and StateFactory in server/patterns/state.js")

    set_p(doc.paragraphs[173], "The base class PatientState defines the contract and throws descriptive domain errors when an invalid transition is attempted.")
    set_p(doc.paragraphs[174], "WaitingState enforces a doctor assignment guard: a patient cannot enter consultation without a valid doctor object.")
    set_p(doc.paragraphs[175], "ConsultationState evaluates severityScore to automatically resolve or escalate clinical alerts to senior specialists.")
    set_p(doc.paragraphs[176], "DischargedState acts as a terminal sink: any call to next() or escalate() throws an error, protecting finalized records.")

    # ---------------------------------------------------------
    # 7. Section 5: System Design & Architecture
    # ---------------------------------------------------------
    print("Updating Section 5 Architecture...")
    set_p(doc.paragraphs[179],
        "BioCare is a three layer application. The browser holds the user interface. An Express server exposes a REST API "
        "and contains the business modules. A data access layer reads and writes MongoDB, or database.json when MongoDB is "
        "not reachable. The three patterns sit in the middle layer, between the route handlers and the data access code: "
        "Factory Method for patient instantiation, Facade for multi-subsystem orchestration, and State for lifecycle governance."
    )
    # P180 is image
    set_caption(doc.paragraphs[181], "Figure 15: ", "Overall architecture and placement of the three patterns (Factory Method, Facade, State Pattern)")

    # Table 7: Layers
    t7 = doc.tables[7]
    set_cell(t7.rows[5].cells[0], "Clinical lifecycle", bold=True)
    set_cell(t7.rows[5].cells[1], "server/patterns/state.js")
    set_cell(t7.rows[5].cells[2], "Patient lifecycle governance, transition guards, and clinical escalation (State).")

    # P191 is image
    set_caption(doc.paragraphs[192], "Figure 16: ", "Sequence of one patient journey: registration, admission and checkout across the three design patterns")

    set_p(doc.paragraphs[195],
        "Consultation. A doctor calls the next patient from the physician terminals. The patient transitions from WaitingState "
        "to ConsultationState via transition guards, and the doctor is assigned."
    )
    set_p(doc.paragraphs[196],
        "Admission. If the doctor decides to admit the patient, AdmissionFacade.admitPatient() coordinates bed allocation, "
        "doctor assignment, initial billing, and executes the state transition to AdmittedState."
    )
    set_p(doc.paragraphs[197],
        "Checkout. AdmissionFacade.dischargePatient() coordinates bed release, doctor release, invoice settlement, and "
        "advances the patient into terminal DischargedState."
    )

    # Table 9: Endpoints
    t9 = doc.tables[9]
    set_cell(t9.rows[6].cells[0], "POST /api/admissions/discharge", bold=True)
    set_cell(t9.rows[6].cells[1], "Check a patient out. Uses AdmissionFacade.dischargePatient() and transitions patient to DischargedState.")

    set_p(doc.paragraphs[202],
        "The patterns do not work in isolation. The patient object created by PatientFactory carries the symptoms and "
        "triage category that initialize the patient into WaitingState. As the patient progresses, the State pattern governs "
        "valid transitions and clinical escalations. When admission or discharge occurs, AdmissionFacade coordinates the bed, "
        "doctor, and billing subsystems while advancing the patient state to AdmittedState and ultimately DischargedState."
    )

    # ---------------------------------------------------------
    # 8. Section 6: Screenshots and Verification
    # ---------------------------------------------------------
    print("Updating Section 6 Screenshots & Tests...")
    set_p(doc.paragraphs[244], "6.6 Patient Lifecycle Transitions and Clinical Governance (State)")
    set_p(doc.paragraphs[245],
        "The clinical lifecycle moves patients through distinct operational states. When a patient arrives at the triage desk, "
        "PatientFactory instantiates the record with an initial WaitingState. Once an available doctor calls the patient, "
        "WaitingState.next(doctor) verifies physician assignment and transitions the patient into ConsultationState."
    )
    # P246 is image
    set_caption(doc.paragraphs[247], "Figure 27: ", "Consultation desk: Doctor assigns consultation and State pattern handles clinical escalation")

    set_p(doc.paragraphs[248],
        "During consultation, clinical escalation rules evaluate vital signs and symptom urgency. High severity patients "
        "(severity score >= 8) are escalated with specialist recommendations and immediate inpatient admission directives. "
        "AdmissionFacade.admitPatient() coordinates bed allocation and transitions the patient into AdmittedState."
    )
    # P249 is image
    set_caption(doc.paragraphs[250], "Figure 28: ", "Inpatient ward admission: AdmissionFacade coordinates subsystems and transitions patient to AdmittedState")

    set_p(doc.paragraphs[251],
        "The State pattern strictly enforces transition guards across all endpoints. Figure 29 shows the terminal execution "
        "log verifying that attempting to enter consultation without a doctor, attempting illegal ward transitions, or "
        "attempting to transition out of terminal DischargedState are immediately caught and rejected with explicit domain errors."
    )
    # P252 is image
    set_caption(doc.paragraphs[253], "Figure 29: ", "Terminal execution log: State pattern transition validation and guard enforcement")

    set_p(doc.paragraphs[254],
        "Finally, checkout is executed via AdmissionFacade.dischargePatient(). The bed is released, the supervising doctor is "
        "freed, the invoice is settled, and the patient transitions safely into DischargedState. The receipt in Figure 30 confirms "
        "the completed checkout and locked terminal state."
    )
    # P255 is image
    set_caption(doc.paragraphs[256], "Figure 30: ", "Completed discharge and settled receipt in DischargedState")

    set_p(doc.paragraphs[258],
        "The repository contains an automated test script (js/test.js) that verifies all three patterns directly. We ran it with "
        "npm test. Figure 31 shows the sections for the three patterns. The factory tests check triage priority assignment and "
        "severity auto-flooring. The state tests check lifecycle transitions, transition guards, clinical escalation, and "
        "terminal state immutability. The facade tests verify coordinated bed allocation, doctor scheduling, and invoice settlement."
    )
    # P259 is image
    set_caption(doc.paragraphs[260], "Figure 31: ", "Output of npm test for the Factory Method, State Pattern, and Facade sections")

    # ---------------------------------------------------------
    # 9. Section 7: Results and Discussion
    # ---------------------------------------------------------
    print("Updating Section 7 Results & Discussion...")
    t10 = doc.tables[10]
    set_cell(t10.rows[8].cells[0], "State transition: Waiting -> Consultation")
    set_cell(t10.rows[8].cells[1], "State")
    set_cell(t10.rows[8].cells[2], "Transition succeeded when doctor assigned; blocked with error if doctor missing.")

    set_cell(t10.rows[9].cells[0], "Clinical escalation during consultation (severity >= 8)")
    set_cell(t10.rows[9].cells[1], "State")
    set_cell(t10.rows[9].cells[2], "Handled urgent escalation: assigned to Senior Attending Specialist with immediate ICU prioritization.")

    # Check if we should update or add rows in t10
    if len(t10.rows) == 11:
        # Add new row for terminal guard check before npm test
        row_new = t10.add_row()
        set_cell(t10.rows[10].cells[0], "Terminal guard check: transition from Discharged")
        set_cell(t10.rows[10].cells[1], "State")
        set_cell(t10.rows[10].cells[2], "Blocked with domain error: Cannot transition from terminal Discharged state.")

        set_cell(row_new.cells[0], "npm test (Full Test Suite)")
        set_cell(row_new.cells[1], "All three")
        set_cell(row_new.cells[2], "All tests for Factory Method, State Pattern, and Facade passed successfully (100%).")

    set_p(doc.paragraphs[267],
        "The main benefit we noticed is that each pattern follows a different direction of change. If the hospital adds a patient "
        "category, only the factory and one subclass change. If a new subsystem is needed during admission, only the facade changes. "
        "If a new clinical state or transition guard is introduced, one new state class derived from PatientState is added without "
        "touching routing controllers."
    )

    set_p(doc.paragraphs[268],
        "Adding a new operational state is a good example. To add a hypothetical 24-hour ObservationState, we write an "
        "ObservationState class extending PatientState, define valid transitions to and from ConsultationState and AdmittedState, "
        "and register it in StateFactory. Neither the Express route handlers nor AdmissionFacade require modification."
    )

    set_p(doc.paragraphs[274],
        "State pattern. State rehydration in the current implementation uses StateFactory.getState() based on persisted state "
        "strings. While this cleanly re-creates state instances, caching or pooling state flyweights could further optimize "
        "high-throughput hospital server deployments."
    )

    # ---------------------------------------------------------
    # 10. Section 8: Conclusion
    # ---------------------------------------------------------
    print("Updating Section 8 Conclusion...")
    set_p(doc.paragraphs[277],
        "BioCare EHR tracks a patient from initial triage to final checkout, structured around three core object-oriented "
        "design patterns: Factory Method (Creational) centralizes domain entity creation and category priority rules; "
        "Facade (Structural) unifies complex multi-subsystem ward admission and discharge operations; and State (Behavioural) "
        "governs clinical lifecycle transitions, transition guards, and triage escalation."
    )

    set_p(doc.paragraphs[278],
        "We verified system correctness through full browser workflows, REST API calls, and the automated test suite (npm test). "
        "All tests passed with 100% success, confirming that triage auto-flooring, state transition guards, and facade "
        "orchestrations operate reliably under real runtime conditions."
    )

    set_p(doc.paragraphs[279],
        "The report also identifies areas for future enhancements: decoupling prescription cost calculation from the factory, "
        "introducing transactional rollbacks in the facade, and pooling state instances. These enhancements can be introduced "
        "within the existing architectural boundaries without modifying other layers."
    )

    print(f"Saving updated document to {output_path}...")
    doc.save(output_path)
    print("Update complete!")

if __name__ == '__main__':
    src = r'C:\Users\saisa\Downloads\BioCare_Design_Patterns_Final_Report_BACKUP.docx'
    dst = r'C:\Users\saisa\Downloads\BioCare_Design_Patterns_Final_Report.docx'
    update_document(src, dst)
